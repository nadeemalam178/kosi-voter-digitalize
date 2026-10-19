# -*- coding: utf-8 -*-
import openpyxl
import re
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

# Target file
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)

EXCEL_PATH = os.path.abspath(os.path.join(SCRIPT_DIR, "..", "Digitalized_Voter_Data.xlsx"))
if not os.path.exists(EXCEL_PATH):
    EXCEL_PATH = "Digitalized_Voter_Data.xlsx"

def clean_phone(phone):
    if not phone:
        return ""
    p = str(phone).strip()
    p = re.sub(r'[^0-9]', '', p)
    if len(p) > 10:
        if p.startswith('91') and len(p) == 12:
            p = p[2:]
        elif len(p) > 10:
            p = p[:10]
    return p

import openpyxl.cell.cell as ccell

def clean_name(name):
    if not name:
        return ""
    s = str(name).strip()
    s = ccell.ILLEGAL_CHARACTERS_RE.sub('', s)
    return s

print(f"Loading {EXCEL_PATH}...")
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active
max_r = ws.max_row
print(f"Initial Excel max row: {max_r}")

last_sno = ws.cell(row=max_r, column=1).value
try:
    current_sno = int(last_sno) + 1
except:
    current_sno = max_r

print(f"Starting Serial Number: {current_sno}")

new_rows = []

def add_voter(district, block, panchayat, village, name, father_husband="", mobile="", roll_no="", notes=""):
    global current_sno
    row = [
        current_sno,
        clean_name(district),
        clean_name(block),
        clean_name(panchayat),
        clean_name(village),
        clean_name(name),
        clean_name(father_husband),
        clean_phone(mobile),
        clean_name(roll_no),
        clean_name(notes)
    ]
    new_rows.append(row)
    current_sno += 1

print("Importing modules and populating voters...")
import voters_part1
import voters_part2
import voters_part3
import voters_part4

voters_part1.populate_group1(add_voter)
print(f"After Part 1: {len(new_rows)} voters collected.")

voters_part2.populate_group2(add_voter)
print(f"After Part 2: {len(new_rows)} voters collected.")

voters_part3.populate_group3(add_voter)
print(f"After Part 3: {len(new_rows)} voters collected.")

voters_part4.populate_group4(add_voter)
print(f"After Part 4: {len(new_rows)} voters collected.")

print(f"Appending {len(new_rows)} new rows to Excel worksheet...")
for r in new_rows:
    ws.append(r)

print(f"Saving workbook to {EXCEL_PATH}...")
wb.save(EXCEL_PATH)
print("Saved successfully!")

# Verify workbook after saving
wb_check = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
ws_check = wb_check.active
print(f"Verification: New max row is {ws_check.max_row} (Old was {max_r}). Added {ws_check.max_row - max_r} rows.")
